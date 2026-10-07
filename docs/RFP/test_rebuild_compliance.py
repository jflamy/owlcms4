import contextlib
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import rebuild_compliance


class RebuildTest(unittest.TestCase):
    def test_existing_section_headings_are_kept_not_recreated(self):
        document = Document()
        document.styles.add_style(rebuild_compliance.REQUIREMENT_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        document.add_heading("A. Company", 1)
        document.add_heading("A.1 Eligibility", 2)
        document.add_paragraph("[RFP §2.1] Be eligible.", style=rebuild_compliance.REQUIREMENT_STYLE)
        document.add_paragraph("We are eligible.")
        document.add_heading("L. Compliance matrix", 1)
        kept = document.add_heading("RFP Section 2", 2)
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), "90")
        start.set(qn("w:name"), "_Toc123")
        kept._element.append(start)
        obsolete = document.add_heading("RFP Section 9", 2)
        document.add_heading("Appendices", 1)

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            output = Path(directory) / "output.docx"
            document.save(source)
            with patch("sys.argv", ["rebuild_compliance.py", str(source), str(output)]), \
                    contextlib.redirect_stdout(io.StringIO()):
                rebuild_compliance.main()
            rebuilt = Document(output)

        labels = [p for p in rebuilt.paragraphs if p.text.startswith("RFP Section")]
        self.assertEqual([p.text for p in labels], ["RFP Section 2"])
        self.assertEqual(
            [b.get(qn("w:name")) for b in labels[0]._element.iter(qn("w:bookmarkStart"))],
            ["_Toc123"],
        )
        self.assertIsNone(rebuilt.settings.element.find(qn("w:updateFields")))
        self.assertNotIn(obsolete.text, [p.text for p in rebuilt.paragraphs])


class VersionSelectionTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.directory = Path(self.temporary_directory.name)

    def create_file(self, name):
        path = self.directory / name
        path.touch()
        return path

    def test_baseline_is_version_zero(self):
        baseline = self.create_file("OTIS_Proposal_IWF.docx")
        self.assertEqual(
            rebuild_compliance.next_version(self.directory),
            (baseline, self.directory / "OTIS_Proposal_01.docx"),
        )

    def test_latest_numbered_proposal_takes_precedence(self):
        self.create_file("OTIS_Proposal_IWF.docx")
        self.create_file("OTIS_Proposal_01.docx")
        latest = self.create_file("OTIS_Proposal_09.docx")
        self.assertEqual(
            rebuild_compliance.next_version(self.directory),
            (latest, self.directory / "OTIS_Proposal_10.docx"),
        )

    def test_explicit_version_zero_takes_precedence_over_baseline(self):
        self.create_file("OTIS_Proposal_IWF.docx")
        zero = self.create_file("OTIS_Proposal_00.docx")
        self.assertEqual(
            rebuild_compliance.next_version(self.directory),
            (zero, self.directory / "OTIS_Proposal_01.docx"),
        )

    def test_versions_continue_beyond_two_digits(self):
        self.create_file("OTIS_Proposal_99.docx")
        latest = self.create_file("OTIS_Proposal_100.docx")
        self.assertEqual(
            rebuild_compliance.next_version(self.directory),
            (latest, self.directory / "OTIS_Proposal_101.docx"),
        )

    def test_missing_proposal_fails_explicitly(self):
        with self.assertRaisesRegex(RuntimeError, "baseline was found"):
            rebuild_compliance.next_version(self.directory)

    def test_default_arguments_use_selected_version(self):
        source = self.create_file("OTIS_Proposal_IWF.docx")
        output = self.directory / "OTIS_Proposal_01.docx"
        with patch("sys.argv", ["rebuild_compliance.py"]), patch.object(
            rebuild_compliance, "next_version", return_value=(source, output)
        ):
            self.assertEqual(rebuild_compliance.parse_arguments(), (source, output))

    def test_explicit_paths_do_not_require_default_proposals(self):
        source = self.create_file("custom.docx")
        output = self.directory / "new.docx"
        with patch("sys.argv", ["rebuild_compliance.py", str(source), str(output)]), patch.object(
            rebuild_compliance, "next_version", side_effect=RuntimeError("No defaults")
        ):
            self.assertEqual(rebuild_compliance.parse_arguments(), (source, output))

    def test_existing_output_is_rejected(self):
        source = self.create_file("source.docx")
        output = self.create_file("existing.docx")
        with patch("sys.argv", ["rebuild_compliance.py", str(source), str(output)]):
            with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as error:
                rebuild_compliance.parse_arguments()
        self.assertEqual(error.exception.code, 2)
        self.assertTrue(source.exists())
        self.assertTrue(output.exists())

    def test_source_cannot_be_overwritten(self):
        source = self.create_file("source.docx")
        with patch("sys.argv", ["rebuild_compliance.py", str(source), str(source)]):
            with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as error:
                rebuild_compliance.parse_arguments()
        self.assertEqual(error.exception.code, 2)
        self.assertTrue(source.exists())


if __name__ == "__main__":
    unittest.main()
