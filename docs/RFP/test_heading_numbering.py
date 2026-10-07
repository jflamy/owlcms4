from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import heading_numbering
import number_headings


def add_bookmark(paragraph, name, identifier):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    heading_numbering.insert_after_pPr(paragraph._element, start)
    paragraph._element.append(end)
    return start


def add_hyperlink(paragraph, anchor, text):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._element.append(hyperlink)


class HeadingNumberingTest(unittest.TestCase):
    def build(self):
        document = Document()
        document.add_heading("Cover Letter", 1)
        document.add_heading("A. Company", 1)
        document.add_heading("A.1 Eligibility", 2)
        self.target = document.add_heading("A.1.1 Minimum requirements", 3)
        document.add_heading("B. Solution", 1)
        document.add_heading("B.1 Overview", 2)
        document.add_heading("Appendices", 1)
        add_bookmark(self.target, "Ref_A_1_1", 1)
        self.reference = document.add_paragraph()
        add_hyperlink(self.reference, "Ref_A_1_1", "A.1.1")
        return document

    def test_conversion_moves_numbers_into_outline_list(self):
        document = self.build()
        report = number_headings.convert(document)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "converted.docx"
            document.save(path)
            number_headings.validate(path, report["expected"])
            converted = Document(path)
        headings = [p for p in converted.paragraphs if p.style.name.startswith("Heading")]
        numbers = heading_numbering.heading_numbers(converted)
        self.assertEqual(
            [(numbers.get(p._element), p.text) for p in headings],
            [
                (None, "Cover Letter"),
                ("A", "Company"),
                ("A.1", "Eligibility"),
                ("A.1.1", "Minimum requirements"),
                ("B", "Solution"),
                ("B.1", "Overview"),
                (None, "Appendices"),
            ],
        )
        self.assertEqual(report["unnumbered"], ["Cover Letter", "Appendices"])

    def test_first_level_is_lettered_and_deeper_levels_are_decimal(self):
        document = self.build()
        number_headings.convert(document)
        numbering = document.part.numbering_part.element
        num_id = heading_numbering._style_num_id(document.styles["Heading 1"])
        num = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(num_id))
        abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        abstract = next(
            a for a in numbering.findall(qn("w:abstractNum"))
            if a.get(qn("w:abstractNumId")) == abstract_id
        )
        levels = abstract.findall(qn("w:lvl"))
        formats = [level.find(qn("w:numFmt")).get(qn("w:val")) for level in levels[:3]]
        texts = [level.find(qn("w:lvlText")).get(qn("w:val")) for level in levels[:3]]
        self.assertEqual(formats, ["upperLetter", "decimal", "decimal"])
        self.assertEqual(texts, ["%1.", "%1.%2", "%1.%2.%3"])

    def test_reference_hyperlink_becomes_heading_number_field(self):
        document = self.build()
        report = number_headings.convert(document)
        self.assertEqual(report["hyperlinks"], 1)
        instructions = [i.text for i in self.reference._element.iter(qn("w:instrText"))]
        self.assertEqual(instructions, [r" REF Ref_A_1_1 \w \h "])
        self.assertEqual(self.reference.text, "A.1.1")
        self.assertEqual(heading_numbering.field_bookmark(self.reference._element), "Ref_A_1_1")

    def test_drifted_bookmark_start_returns_to_heading(self):
        document = Document()
        heading = document.add_heading("A. Company", 1)
        body_text = document.add_paragraph("Text")
        start = add_bookmark(heading, "Ref_A", 7)
        start.getparent().remove(start)
        body_text._element.append(start)
        self.assertEqual(heading_numbering.repair_heading_bookmarks(document), ["Ref_A"])
        self.assertIs(start.getparent(), heading._element)
        self.assertIs(start.getprevious(), heading._element.pPr)

    def test_gap_in_typed_numbers_is_rejected(self):
        document = Document()
        document.add_heading("A. Company", 1)
        document.add_heading("A.2 Skipped", 2)
        with self.assertRaisesRegex(RuntimeError, "renumbered A.1"):
            heading_numbering.check_typed_sequence(document)

    def test_styles_derived_from_headings_are_not_numbered(self):
        document = self.build()
        derived = document.styles.add_style("Derived Heading", WD_STYLE_TYPE.PARAGRAPH)
        derived.base_style = document.styles["Heading 1"]
        number_headings.convert(document)
        numPr = derived.element.pPr.numPr
        self.assertEqual(numPr.numId.val, 0)
        self.assertIsNone(numPr.ilvl)

    def test_conversion_refuses_to_run_twice(self):
        document = self.build()
        number_headings.convert(document)
        with self.assertRaisesRegex(RuntimeError, "already"):
            number_headings.convert(document)


if __name__ == "__main__":
    unittest.main()
