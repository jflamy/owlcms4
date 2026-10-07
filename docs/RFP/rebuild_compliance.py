#!/usr/bin/env python3

import argparse
from collections import defaultdict
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

import heading_numbering


TABLE_HEADERS = ("Requirement", "Compliance", "Reference", "Justification")
REQUIREMENT_STYLE = "requirementNarrative"
RFP_PATTERN = re.compile(r"^\[RFP §([^]]+)]\s*(.+)$")


def clean(text):
    return " ".join(text.split())


def next_version(directory):
    versions = []
    for path in directory.glob("OTIS_Proposal_*.docx"):
        match = re.fullmatch(r"OTIS_Proposal_(\d+)\.docx", path.name)
        if match:
            versions.append((int(match.group(1)), path))
    if not versions:
        baseline = directory / "OTIS_Proposal_IWF.docx"
        if not baseline.is_file():
            raise RuntimeError(
                "No numbered OTIS proposal or OTIS_Proposal_IWF.docx baseline was found"
            )
        versions.append((0, baseline))
    number, source = max(versions)
    return source, directory / f"OTIS_Proposal_{number + 1:02d}.docx"


def parse_arguments(
    description="Rebuild RFP compliance tables from requirementNarrative paragraphs.",
):
    parser = argparse.ArgumentParser(
        description=description,
        epilog=(
            "Defaults to the latest numbered proposal beside this script, or "
            "OTIS_Proposal_IWF.docx as version 00, and writes the next numbered file "
            "without overwriting an existing file."
        ),
    )
    parser.add_argument("source", nargs="?", type=Path)
    parser.add_argument("output", nargs="?", type=Path)
    arguments = parser.parse_args()
    directory = Path(__file__).resolve().parent
    source = arguments.source
    output = arguments.output
    if source is None or output is None:
        default_source, default_output = next_version(directory)
        source = source or default_source
        output = output or default_output
    if source.resolve() == output.resolve():
        parser.error("Source and output must be different files")
    if output.exists():
        parser.error(f"Output already exists: {output}")
    return source, output


def is_compliance_table(table):
    return tuple(cell.text for cell in table.rows[0].cells) == TABLE_HEADERS


def heading_before(paragraph, document):
    element = paragraph._element.getprevious()
    while element is not None:
        if element.tag.endswith("}p"):
            candidate = Paragraph(element, document)
            if candidate.style.name.startswith("Heading "):
                return candidate
        element = element.getprevious()
    raise RuntimeError(f"No heading precedes requirement: {paragraph.text}")


def justification_after(paragraph, document):
    lines = []
    element = paragraph._element.getnext()
    while element is not None:
        if element.tag.endswith("}tbl"):
            break
        if element.tag.endswith("}p"):
            candidate = Paragraph(element, document)
            if candidate.style.name.startswith("Heading "):
                break
            if candidate.style.name == REQUIREMENT_STYLE:
                break
            text = clean(candidate.text)
            if text:
                lines.append(text)
        element = element.getnext()
    return "\n".join(lines) or "[TO BE COMPLETED]"


def bookmark_for_heading(heading, reference, document, next_bookmark_id):
    bookmarks = heading._element.xpath(".//w:bookmarkStart")
    if bookmarks:
        return bookmarks[0].get(qn("w:name")), next_bookmark_id

    bookmark_name = heading_numbering.unique_bookmark_name(
        document, "Ref_" + re.sub(r"[^A-Za-z0-9_]", "_", reference)
    )
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(next_bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(next_bookmark_id))
    heading_numbering.insert_after_pPr(heading._element, start)
    heading._element.append(end)
    return bookmark_name, next_bookmark_id + 1


def add_internal_link(paragraph, display_text, bookmark_name):
    """Heading-number cross-reference (REF \\w \\h) that Word keeps up to date."""
    heading_numbering.append_heading_number_ref(
        paragraph._element, bookmark_name, display_text
    )


def set_repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_column_widths(table, document):
    section = document.sections[0]
    available_width = section.page_width - section.left_margin - section.right_margin
    widths = [
        int(available_width * proportion)
        for proportion in (0.22, 0.14, 0.12, 0.52)
    ]
    table.autofit = False
    layout = table._tbl.tblPr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table._tbl.tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for index, width in enumerate(widths):
        table.columns[index].width = width
        table._tbl.tblGrid.gridCol_lst[index].w = width
        for cell in table.columns[index].cells:
            cell.width = width


def origin_sort_key(origin):
    match = re.match(r"^(\d+)(?:\.(.+))?$", origin)
    if not match:
        return 999, ((999, origin),)
    section = int(match.group(1))
    suffix = match.group(2)
    if suffix is None:
        return section, ()
    parts = []
    for part in suffix.split("."):
        if part.isdigit():
            parts.append((1, int(part)))
        elif len(part) == 1 and part.isalpha():
            parts.append((0, ord(part.upper()) - ord("A")))
        else:
            parts.append((2, part))
    return section, tuple(parts)


def section_for_origin(origin):
    match = re.match(r"^(\d+)", origin)
    if not match:
        raise RuntimeError(f"Cannot determine RFP section from origin: {origin}")
    return int(match.group(1))


def generated_requirement_id(record, used_ids, origin_counts):
    heading_text = clean(record["heading"].text)
    gate = re.search(r"\bHG-\d{2}\b", heading_text)
    if gate and gate.group(0) not in used_ids:
        return gate.group(0)
    profile = re.search(r"\b(?:M0|M1|M2|M3|MP)\b", heading_text)
    if profile:
        candidate = f"{record['origin']}-{profile.group(0)}"
        if candidate not in used_ids:
            return candidate
    origin_counts[record["origin"]] += 1
    candidate = record["origin"]
    if candidate not in used_ids:
        return candidate
    return f"{record['origin']}-R{origin_counts[record['origin']]:02d}"


def validate_output(path, expected_requirements, expected_sections):
    document = Document(path)
    requirements = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == REQUIREMENT_STYLE
    ]
    tables = [table for table in document.tables if is_compliance_table(table)]
    rows = [row for table in tables for row in table.rows[1:]]
    if len(requirements) != expected_requirements or len(rows) != expected_requirements:
        raise RuntimeError(
            "Generated requirement/table count mismatch: "
            f"{len(requirements)} paragraphs, {len(rows)} rows, "
            f"expected {expected_requirements}"
        )
    if len(tables) != expected_sections:
        raise RuntimeError(
            f"Generated {len(tables)} section tables; expected {expected_sections}"
        )

    bookmark_names = {
        bookmark.get(qn("w:name"))
        for bookmark in document.element.xpath(".//w:bookmarkStart")
    }
    references = []
    for table in tables:
        properties = table.rows[0]._tr.trPr
        if properties is None or properties.find(qn("w:tblHeader")) is None:
            raise RuntimeError("A compliance table header is not set to repeat")
        for row in table.rows[1:]:
            if row.cells[1].text != "Compliant":
                raise RuntimeError("A compliance cell does not contain 'Compliant'")
            anchor = heading_numbering.field_bookmark(row.cells[2]._element)
            if anchor is None:
                raise RuntimeError(
                    f"Reference {row.cells[2].text!r} does not contain a cross-reference"
                )
            if anchor not in bookmark_names:
                raise RuntimeError(f"Reference points to missing bookmark: {anchor}")
            references.append(clean(row.cells[2].text))
    if len(references) != len(set(references)):
        raise RuntimeError("Duplicate proposal references exist in the compliance tables")


def main():
    source, output = parse_arguments()
    document = Document(source)
    body = document.element.body

    if REQUIREMENT_STYLE not in document.styles:
        raise RuntimeError(f"Paragraph style not found: {REQUIREMENT_STYLE}")

    existing_ids = {}
    compliance_tables = []
    for table in document.tables:
        if not is_compliance_table(table):
            continue
        compliance_tables.append(table)
        for row in table.rows[1:]:
            reference = heading_numbering.field_bookmark(row.cells[2]._element)
            reference = reference or clean(row.cells[2].text)
            requirement_id = clean(row.cells[0].text).split()[0]
            existing_ids[reference] = requirement_id

    numbers = heading_numbering.heading_numbers(document)
    requirements = []
    max_bookmark_id = max(
        (int(bookmark.get(qn("w:id"))) for bookmark in document.element.xpath(".//w:bookmarkStart")),
        default=0,
    )
    next_bookmark_id = max_bookmark_id + 1
    for document_order, paragraph in enumerate(document.paragraphs):
        if paragraph.style.name != REQUIREMENT_STYLE:
            continue
        match = RFP_PATTERN.match(clean(paragraph.text))
        if not match:
            raise RuntimeError(
                f"{REQUIREMENT_STYLE} paragraph lacks an [RFP §…] prefix: {paragraph.text}"
            )
        origin, requirement_text = match.groups()
        heading = heading_before(paragraph, document)
        reference = numbers.get(heading._element)
        if reference is None:
            raise RuntimeError(f"Requirement follows an unnumbered heading: {heading.text}")
        bookmark_name, next_bookmark_id = bookmark_for_heading(
            heading, reference, document, next_bookmark_id
        )
        requirements.append({
            "origin": origin,
            "requirement": requirement_text,
            "heading": heading,
            "reference": reference,
            "bookmark": bookmark_name,
            "justification": justification_after(paragraph, document),
            "document_order": document_order,
        })

    if not requirements:
        raise RuntimeError(f"No {REQUIREMENT_STYLE} paragraphs were found")

    used_ids = set(existing_ids.values())
    origin_counts = defaultdict(int)
    for record in requirements:
        requirement_id = existing_ids.get(record["bookmark"])
        if requirement_id is None:
            requirement_id = existing_ids.get(record["reference"])
        if requirement_id is None:
            requirement_id = generated_requirement_id(record, used_ids, origin_counts)
        record["requirement_id"] = requirement_id
        used_ids.add(requirement_id)

    requirements.sort(
        key=lambda record: (origin_sort_key(record["origin"]), record["document_order"])
    )

    for table in compliance_tables:
        body.remove(table._element)
    section_headings = {}
    for paragraph in document.paragraphs:
        match = re.fullmatch(r"RFP Section (\d+)", paragraph.text)
        if match:
            if int(match.group(1)) in section_headings:
                body.remove(paragraph._element)
            else:
                section_headings[int(match.group(1))] = paragraph

    appendices = next(
        paragraph for paragraph in document.paragraphs if paragraph.text == "Appendices"
    )
    by_section = defaultdict(list)
    for record in requirements:
        by_section[section_for_origin(record["origin"])].append(record)

    for section in set(section_headings) - set(by_section):
        body.remove(section_headings.pop(section)._element)

    for section_index, section in enumerate(sorted(by_section)):
        heading = section_headings.get(section)
        if heading is None:
            heading = document.add_paragraph(f"RFP Section {section}", style="Heading 2")
        heading_numbering.suppress_numbering(heading)
        heading.paragraph_format.page_break_before = section_index > 0
        appendices._element.addprevious(heading._element)

        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        set_column_widths(table, document)
        for cell, header in zip(table.rows[0].cells, TABLE_HEADERS):
            cell.text = header
            for run in cell.paragraphs[0].runs:
                run.bold = True
        set_repeat_header(table.rows[0])

        for record in by_section[section]:
            row = table.add_row()
            row.cells[0].text = (
                f"{record['requirement_id']}  {record['requirement']}"
            )
            row.cells[1].text = "Compliant"
            add_internal_link(
                row.cells[2].paragraphs[0],
                record["reference"],
                record["bookmark"],
            )
            row.cells[3].text = record["justification"]
        appendices._element.addprevious(table._element)

    document.save(output)
    validate_output(output, len(requirements), len(by_section))
    print(f"Source: {source}")
    print(f"Wrote: {output}")
    print(f"Requirements/rows: {len(requirements)}")
    print(f"RFP section tables: {len(by_section)}")
    print("Validation: passed")


if __name__ == "__main__":
    main()