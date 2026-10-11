import contextlib
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

import compliance_summaries
import rebuild_compliance


def build_source(path, justification="We are eligible because of many verbose reasons."):
    document = Document()
    document.styles.add_style(rebuild_compliance.REQUIREMENT_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    document.add_heading("A. Company", 1)
    document.add_heading("A.1 Eligibility", 2)
    document.add_paragraph(
        "[RFP §2.1] Be eligible in a long-winded way.", style=rebuild_compliance.REQUIREMENT_STYLE
    )
    document.add_paragraph(justification)
    document.add_heading("Appendices", 1)
    document.save(path)


class SummariesRoundTripTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.directory = Path(self.temporary_directory.name)
        self.source = self.directory / "source.docx"
        self.markdown = self.directory / "summaries.md"
        build_source(self.source)

    def run_quietly(self, function, *arguments):
        with contextlib.redirect_stdout(io.StringIO()) as output:
            function(*arguments)
        return output.getvalue()

    def rebuild(self, source):
        output = self.directory / "output.docx"
        output.unlink(missing_ok=True)
        output.with_suffix(".md").unlink(missing_ok=True)
        argv = ["rebuild_compliance.py", str(source), str(output), "--summaries", str(self.markdown)]
        with patch("sys.argv", argv):
            log = self.run_quietly(rebuild_compliance.main)
        table = next(t for t in Document(output).tables if rebuild_compliance.is_compliance_table(t))
        return [cell.text for cell in table.rows[1].cells], log

    def test_summaries_replace_verbose_text(self):
        self.run_quietly(compliance_summaries.extract, self.source, self.markdown)
        self.run_quietly(
            compliance_summaries.fill, self.markdown, "2.1 | Be eligible. | Eligible.\n"
        )
        cells, log = self.rebuild(self.source)
        self.assertEqual(cells[0], "2.1  Be eligible.")
        self.assertEqual(cells[3], "Eligible.")
        self.assertNotIn("Warning", log)

    def test_changed_source_keeps_original_and_reextract_marks_stale(self):
        self.run_quietly(compliance_summaries.extract, self.source, self.markdown)
        self.run_quietly(
            compliance_summaries.fill, self.markdown, "2.1 | Be eligible. | Eligible.\n"
        )
        changed = self.directory / "changed.docx"
        build_source(changed, justification="New wording.")
        cells, log = self.rebuild(changed)
        self.assertEqual(cells[3], "New wording.")
        self.assertIn("source text changed", log)

        log = self.run_quietly(compliance_summaries.extract, changed, self.markdown)
        self.assertIn("1 stale", log)
        self.assertIn("<!-- old J: Eligible. -->", self.markdown.read_text())

    def test_rebuild_writes_md_and_word_edits_are_read_back(self):
        self.run_quietly(compliance_summaries.extract, self.source, self.markdown)
        self.run_quietly(
            compliance_summaries.fill, self.markdown, "2.1 | Be eligible. | Eligible. | Typo.\n"
        )
        self.rebuild(self.source)
        output = self.directory / "output.docx"
        sidecar = compliance_summaries.read_summaries(output.with_suffix(".md"))["2.1"]
        self.assertEqual(
            (sidecar["requirement"], sidecar["justification"], sidecar["fix"]),
            ("Be eligible.", "Eligible.", "Typo."),
        )

        document = Document(output)
        table = next(t for t in document.tables if rebuild_compliance.is_compliance_table(t))
        table.rows[1].cells[3].text = "Edited in Word."
        document.save(output)
        self.run_quietly(
            compliance_summaries.extract, output, output.with_suffix(".md"), None, True
        )
        sidecar = compliance_summaries.read_summaries(output.with_suffix(".md"))["2.1"]
        self.assertEqual(sidecar["justification"], "Edited in Word.")
        self.assertEqual(sidecar["fix"], "Typo.")

    def test_fill_rejects_unknown_ids(self):
        self.run_quietly(compliance_summaries.extract, self.source, self.markdown)
        with self.assertRaisesRegex(RuntimeError, "Unknown requirement IDs: 9.9"):
            compliance_summaries.fill(self.markdown, "9.9 | x | y\n")

    def test_fix_flag_persists_and_warns(self):
        self.run_quietly(compliance_summaries.extract, self.source, self.markdown)
        self.run_quietly(
            compliance_summaries.fill, self.markdown, "2.1 | Be eligible. | = | Typo here.\n"
        )
        self.run_quietly(compliance_summaries.extract, self.source, self.markdown)
        text = self.markdown.read_text()
        self.assertIn("- ⚠️ **2.1**: Typo here.", text)
        self.assertIn("FIX: Typo here.", text)
        _, log = self.rebuild(self.source)
        self.assertIn("2.1: FIX in proposal: Typo here.", log)

        self.run_quietly(compliance_summaries.fill, self.markdown, "2.1 | = | = | -\n")
        self.assertNotIn("\nFIX:", self.markdown.read_text())


if __name__ == "__main__":
    unittest.main()
