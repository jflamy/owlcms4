#!/usr/bin/env python3
"""Convert typed heading prefixes ("A.1.1 Title") into Word outline numbering.

After conversion, Word sees the prefixes as real heading numbers, so
Insert > Cross-reference > Heading number works and a generated table of
contents shows the numbers. Compliance-table references become REF fields
that follow the numbering when fields are updated.
"""

import sys

from docx import Document
from docx.oxml.ns import qn

import heading_numbering as numbering
import rebuild_compliance


def restyle_figure_headings(document):
    """Image-only paragraphs that were styled as headings by accident."""
    restyled = 0
    for paragraph in document.paragraphs:
        if numbering.heading_level(paragraph) is None or paragraph.text.strip():
            continue
        if paragraph._element.find(".//" + qn("w:drawing")) is None:
            continue
        paragraph.style = document.styles["Normal"]
        restyled += 1
    return restyled


def convert(document):
    if numbering._style_num_id(document.styles["Heading 1"]) is not None:
        raise RuntimeError("Heading styles already carry outline numbering")

    repaired = numbering.repair_heading_bookmarks(document)
    restyled = restyle_figure_headings(document)
    numbering.check_typed_sequence(document)

    numbers = numbering.heading_numbers(document)
    bookmarks = numbering.heading_bookmarks(document)
    hyperlinks = numbering.convert_number_hyperlinks(document, bookmarks, numbers)
    refs = numbering.split_heading_text_refs(document, bookmarks, numbers)

    headings = [
        paragraph
        for paragraph in document.paragraphs
        if numbering.heading_level(paragraph) is not None
    ]
    numbered = {paragraph._element for paragraph in headings if paragraph._element in numbers}
    numbering.ensure_heading_list(document)
    unnumbered = []
    for paragraph in headings:
        if paragraph._element in numbered:
            numbering.strip_typed_prefix(paragraph)
        else:
            numbering.suppress_numbering(paragraph)
            unnumbered.append(paragraph.text)

    return {
        "expected": sorted(
            (numbers[p._element], p.text) for p in headings if p._element in numbered
        ),
        "repaired": repaired,
        "restyled": restyled,
        "hyperlinks": hyperlinks,
        "refs": refs,
        "unnumbered": unnumbered,
    }


def validate(path, expected):
    document = Document(path)
    numbers = numbering.heading_numbers(document)
    actual = sorted(
        (numbers[p._element], p.text) for p in document.paragraphs if p._element in numbers
    )
    if actual != expected:
        raise RuntimeError("Outline numbering does not reproduce the typed heading numbers")
    for paragraph in document.paragraphs:
        if paragraph._element in numbers and numbering.TYPED_PREFIX.match(paragraph.text):
            raise RuntimeError(f"Typed prefix remains on numbered heading: {paragraph.text}")

    body = document.element.body
    names = {start.get(qn("w:name")) for start in body.iter(qn("w:bookmarkStart"))}
    for instruction in body.iter(qn("w:instrText")):
        match = numbering.REF_INSTRUCTION.match(instruction.text or "")
        if match and match.group(1) not in names:
            raise RuntimeError(f"REF field points to missing bookmark: {match.group(1)}")
    for hyperlink in body.iter(qn("w:hyperlink")):
        anchor = hyperlink.get(qn("w:anchor"))
        if anchor and hyperlink.get(qn("r:id")) is None and anchor not in names:
            raise RuntimeError(f"Hyperlink points to missing bookmark: {anchor}")
    return len(actual)


def main():
    source, output = rebuild_compliance.parse_arguments(
        "Convert typed heading prefixes into Word outline numbering."
    )
    document = Document(source)
    report = convert(document)
    document.save(output)
    count = validate(output, report["expected"])
    print(f"Source: {source}")
    print(f"Wrote: {output}")
    print(f"Numbered headings: {count}")
    print(f"Unnumbered headings: {', '.join(report['unnumbered'])}")
    print(f"Bookmarks moved back onto headings: {', '.join(report['repaired']) or 'none'}")
    print(f"Image paragraphs restyled from Heading to Normal: {report['restyled']}")
    print(f"Reference hyperlinks converted to REF fields: {report['hyperlinks']}")
    print(f"Heading REF fields split into number + text: {report['refs']}")
    print("Validation: passed")
    print("In Word: Ctrl+A then F9 (Cmd+A, Fn+F9 on macOS) refreshes field results.")


if __name__ == "__main__":
    sys.exit(main())
