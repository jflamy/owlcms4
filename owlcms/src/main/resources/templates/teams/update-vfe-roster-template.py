import argparse
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment


def clean_signature_spacer(sheet):
    changed = False
    for column in range(1, 11):
        source = sheet.cell(13, column)
        target = sheet.cell(14, column)
        if target._style != source._style:
            target._style = copy(source._style)
            changed = True
    return changed


def check(file):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook.active
    assert sheet["A1"].value == "${vfeReportMarker}", file
    assert 'items="vfeTeams"' in sheet["A1"].comment.text, file
    assert 'multisheet="team.sheetName"' in sheet["A1"].comment.text, file
    assert 'items="team.blocks"' in sheet["A10"].comment.text, file
    assert 'items="block.members"' in sheet["A11"].comment.text, file
    assert 'items="block.violations"' in sheet["A12"].comment.text, file
    assert sheet["J11"].value == "${member.key}", file
    assert sheet["A12"].value == "VFE_WARNING:${violation}", file
    assert sheet["A15"].value == '${t.get("Jury.Signature")}', file
    assert "A15:C15" in {str(region) for region in sheet.merged_cells.ranges}, file
    assert sheet.row_dimensions[1].hidden, file
    assert all(sheet.cell(14, c)._style == sheet.cell(13, c)._style for c in range(1, 11)), file
    assert not any("teamAgeGroupsAsString" in str(cell.value) for row in sheet for cell in row), file
    assert not any(cell.value == " " for row in sheet for cell in row), file
    return workbook


def update(file):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook.active
    if sheet["A1"].value == "${vfeReportMarker}":
        if clean_signature_spacer(sheet):
            temporary = file.with_suffix(".tmp.xlsx")
            workbook.save(temporary)
            check(temporary)
            temporary.replace(file)
        check(file)
        return
    assert sheet["A5"].value == "VFE", file
    assert sheet["J11"].value == "${l.teamAgeGroupsAsString}", file
    assert sheet["J14"].value == "${l.teamAgeGroupsAsString}", file
    assert sheet["A16"].value == '${t.get("Jury.Signature")}', file
    original_print_settings = copy(sheet.page_setup)
    signature_style = copy(sheet["A16"]._style)
    signature_height = sheet.row_dimensions[16].height
    sheet.unmerge_cells("A16:C16")
    for row in sheet.iter_rows(min_row=10, max_row=16):
        for cell in row:
            cell.value = None
            cell.comment = None
    for row in sheet:
        for cell in row:
            if cell.value == " ":
                cell.value = None
    sheet["A1"] = "${vfeReportMarker}"
    sheet.row_dimensions[1].hidden = True
    sheet["A1"].comment = Comment(
        'jx:area(lastCell="AA15")\n'
        'jx:each(items="vfeTeams" var="team" multisheet="team.sheetName" lastCell="J15")',
        "owlcms",
    )
    sheet["A7"] = "${team.name}"
    sheet["A5"] = '${t.get("VFE")}'
    sheet["B8"] = "${vfeNumberHeader}"
    sheet["A10"] = "${block.title}"
    sheet["A10"].comment = Comment(
        'jx:each(items="team.blocks" var="block" lastCell="J13")', "owlcms"
    )
    values = {
        "A11": "${member.athlete.gender}",
        "B11": "${memberIndex+1}",
        "C11": "${member.athlete.lastName}",
        "D11": "${member.athlete.firstName}",
        "E11": "${member.athlete.fullBirthDate}",
        "F11": "${member.athlete.categoryName}",
        "H11": "${member.athlete.entryTotal}",
        "J11": "${member.key}",
        "A12": "VFE_WARNING:${violation}",
        "A15": '${t.get("Jury.Signature")}',
    }
    for coordinate, value in values.items():
        sheet[coordinate] = value
    sheet["A11"].comment = Comment(
        'jx:each(items="block.members" var="member" varIndex="memberIndex" lastCell="J11")', "owlcms"
    )
    sheet["A12"].comment = Comment(
        'jx:each(items="block.violations" var="violation" lastCell="J12")', "owlcms"
    )
    sheet["A15"]._style = signature_style
    sheet.row_dimensions[15].height = signature_height
    sheet.row_dimensions[12].height = 30
    sheet.row_dimensions[13].height = 16.5
    sheet.row_dimensions[16].height = 16.5
    sheet.merge_cells("A15:C15")
    clean_signature_spacer(sheet)
    sheet.page_setup = original_print_settings
    temporary = file.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    check(temporary)
    temporary.replace(file)


parser = argparse.ArgumentParser()
parser.add_argument("--check", action="store_true")
arguments = parser.parse_args()
for name in ("VFE_Teams-A4.xlsx", "VFE_Teams-LETTER.xlsx"):
    template = Path(__file__).parent / name
    if arguments.check:
        check(template)
    else:
        update(template)
    print(f"Validated {template.name}")
