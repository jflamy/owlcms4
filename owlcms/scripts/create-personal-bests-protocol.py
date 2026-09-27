"""Create expanded session protocol templates with athlete personal bests."""

import argparse
import copy
import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter


SOURCE_TO_TARGET = {
    "Protocol-A4.xlsx": "Protocol-PersonalBests-A4.xlsx",
    "Protocol-LETTER.xlsx": "Protocol-PersonalBests-LETTER.xlsx",
}
INSERTED_COLUMNS = (15, 20, 22)
SOURCE_SHIFT_POINTS = (15, 19, 20)
EXPECTED_SOURCE_VALUES = {
    "N10": "${l.bestSnatch}",
    "R10": "${l.bestCleanJerk}",
    "S10": "${l.total}",
}


def remap_column(column):
    return column + sum(shift_point <= column for shift_point in SOURCE_SHIFT_POINTS)


def remap_range(cell_range):
    min_column = remap_column(cell_range.min_col)
    max_column = remap_column(cell_range.max_col)
    return (
        f"{get_column_letter(min_column)}{cell_range.min_row}:"
        f"{get_column_letter(max_column)}{cell_range.max_row}"
    )


def copy_cell_style(source, target):
    target._style = copy.copy(source._style)
    target.number_format = source.number_format
    target.alignment = copy.copy(source.alignment)
    target.protection = copy.copy(source.protection)


def highlighted_value(result, personal_best):
    return (
        "${"
        f"{result} != null && {personal_best} != null && {personal_best} > 0 "
        f"&& {result} > {personal_best} ? \"**\" + {result} + \"**\" : {result}"
        "}"
    )


def validate_source(sheet, path):
    mismatches = []
    for coordinate, expected in EXPECTED_SOURCE_VALUES.items():
        actual = sheet[coordinate].value
        if actual != expected:
            mismatches.append(f"{coordinate}: expected {expected!r}, found {actual!r}")
    comment = sheet["A1"].comment
    if comment is None or comment.text != 'jx:area(lastCell="V23")':
        mismatches.append("A1: expected the current JXLS3 V23 area command")
    if mismatches:
        raise ValueError(f"Unexpected source layout in {path}:\n  " + "\n  ".join(mismatches))


def update_commands(sheet):
    expected_commands = {
        "A1": 'jx:area(lastCell="Y23")',
        "A8": (
            'jx:each(items="athletes" var="group" groupBy="group.categorySortCode" '
            'groupOrder="ASC" lastCell="Y10")'
        ),
        "A10": 'jx:each(items="group.items" var="l"  varIndex="lifterLoop" lastCell="Y10")',
        "A12": 'jx:if(condition="session.records.size() > 0"  lastCell="Y15")',
        "A15": 'jx:each(items="session.records" var="r" lastCell="Y15")',
    }
    for coordinate, text in expected_commands.items():
        sheet[coordinate].comment = Comment(text, "owlcms")


def create_template(source_path, target_path):
    workbook = load_workbook(source_path)
    if "Results" not in workbook.sheetnames:
        raise ValueError(f"Results sheet not found in {source_path}")
    sheet = workbook["Results"]
    validate_source(sheet, source_path)

    merged_ranges = list(sheet.merged_cells.ranges)
    for merged_range in merged_ranges:
        sheet.unmerge_cells(str(merged_range))

    for column in (20, 19, 15):
        sheet.insert_cols(column)

    for merged_range in merged_ranges:
        sheet.merge_cells(remap_range(merged_range))

    for row in range(1, sheet.max_row + 1):
        copy_cell_style(sheet.cell(row, 14), sheet.cell(row, 15))
        copy_cell_style(sheet.cell(row, 19), sheet.cell(row, 20))
        copy_cell_style(sheet.cell(row, 21), sheet.cell(row, 22))

    for column in INSERTED_COLUMNS:
        sheet.column_dimensions[get_column_letter(column)].width = 7.7109375
    sheet.column_dimensions["Y"].width = 12.0

    for cell_range in ("K6:O6", "P6:T6", "U6:V6", "W6:W7", "X6:Y7"):
        for merged_range in list(sheet.merged_cells.ranges):
            if merged_range.min_row <= 7 and merged_range.max_row >= 6 and not (
                merged_range.max_col < sheet[cell_range.split(":")[0]].column
                or merged_range.min_col > sheet[cell_range.split(":")[1]].column
            ):
                sheet.unmerge_cells(str(merged_range))
        sheet.merge_cells(cell_range)

    sheet["K6"] = '${t.get("Results.Snatch")}'
    sheet["P6"] = '${t.get("Results.Clean_and_Jerk")}'
    sheet["U6"] = '${t.get("Results.Total")}'
    sheet["W6"] = '${t.get("Results.Rank")}'
    sheet["X6"] = "${bestRankingTitle}"
    sheet["N7"] = '${t.get("Results.Best")}'
    sheet["O7"] = "${'PB'}"
    sheet["S7"] = '${t.get("Results.Best")}'
    sheet["T7"] = "${'PB'}"
    sheet["U7"] = '${t.get("Results.Total")}'
    sheet["V7"] = "${'PB'}"
    copy_cell_style(sheet["V7"], sheet["U7"])
    sheet["V7"].alignment = copy.copy(sheet["T7"].alignment)

    sheet["N10"] = highlighted_value("l.bestSnatch", "l.personalBestSnatch")
    sheet["O10"] = "${l.personalBestSnatch}"
    sheet["S10"] = highlighted_value("l.bestCleanJerk", "l.personalBestCleanJerk")
    sheet["T10"] = "${l.personalBestCleanJerk}"
    sheet["U10"] = highlighted_value("l.total", "l.personalBestTotal")
    sheet["V10"] = "${l.personalBestTotal}"
    sheet["W10"] = '${l.totalRank < 0 ? "inv." : (l.totalRank > 0 ? l.totalRank : "" )}'
    sheet["X10"] = "${l.bestLifterScore}"

    update_commands(sheet)
    sheet.delete_cols(26, sheet.max_column - 25)
    sheet.print_area = "A1:Y125"

    with tempfile.NamedTemporaryFile(dir=target_path.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        validate_template(temporary_path)
        os.replace(temporary_path, target_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def validate_template(path):
    workbook = load_workbook(path)
    sheet = workbook["Results"]
    expected_values = {
        "O7": "${'PB'}",
        "T7": "${'PB'}",
        "V7": "${'PB'}",
        "O10": "${l.personalBestSnatch}",
        "T10": "${l.personalBestCleanJerk}",
        "V10": "${l.personalBestTotal}",
    }
    errors = []
    for coordinate, expected in expected_values.items():
        if sheet[coordinate].value != expected:
            errors.append(f"{coordinate}: expected {expected!r}, found {sheet[coordinate].value!r}")
    for coordinate in ("N10", "S10", "U10"):
        if "**" not in str(sheet[coordinate].value):
            errors.append(f"{coordinate}: improved-result marker is missing")
    expected_comments = {"A1": "Y23", "A8": "Y10", "A10": "Y10", "A12": "Y15", "A15": "Y15"}
    for coordinate, last_cell in expected_comments.items():
        comment = sheet[coordinate].comment
        if comment is None or f'lastCell="{last_cell}"' not in comment.text:
            errors.append(f"{coordinate}: expected lastCell={last_cell}")
    if sheet.max_column > 25:
        errors.append(f"unexpected cells beyond column Y: {sheet.max_column}")
    if set(INSERTED_COLUMNS) != {sheet["O1"].column, sheet["T1"].column, sheet["V1"].column}:
        errors.append("personal-best columns are not O, T, and V")
    if sheet.print_area != "'Results'!$A$1:$Y$125":
        errors.append(f"print area: expected A1:Y125, found {sheet.print_area!r}")
    if sheet.column_dimensions["Y"].width != 12.0:
        errors.append("column Y must be wide enough for the scoring header")
    if sheet["U7"].border.left.style != "thin" or sheet["U7"].border.right.style != "thin":
        errors.append("Total header must have left and right borders")
    if sheet["V7"].alignment.horizontal != "center":
        errors.append("Total PB header must be centered")
    if errors:
        raise ValueError(f"Invalid personal-best protocol template {path}:\n  " + "\n  ".join(errors))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Validate generated templates without modifying them")
    args = parser.parse_args()

    template_directory = Path(__file__).resolve().parent.parent / "src/main/resources/templates/protocol"
    for source_name, target_name in SOURCE_TO_TARGET.items():
        target_path = template_directory / target_name
        if args.check:
            validate_template(target_path)
            print(f"Checked {target_path}")
        else:
            create_template(template_directory / source_name, target_path)
            print(f"Created {target_path}")


if __name__ == "__main__":
    main()