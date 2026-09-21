#!/usr/bin/env python3
"""Convert weigh-in templates from JXLS1 rows to JXLS3 comments."""

import argparse
import copy
import os
from pathlib import Path
import tempfile

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Border, Side
from openpyxl.utils.cell import range_boundaries


TEMPLATE_DIRECTORY = (
    Path(__file__).resolve().parents[1] / "src/main/resources/templates/weighin"
)
TEMPLATE_NAMES = (
    "WeighInForm-A4.xlsx",
    "WeighInForm-LETTER.xlsx",
    "WeighInForm_Horizontal-A4.xlsx",
    "WeighInForm_Horizontal-LETTER.xlsx",
    "_WeighInForm_x2-A4.xlsx",
    "_WeighInForm_x2-LETTER.xlsx",
    "_WeighInForm_x2_Horizontal-A4.xlsx",
    "_WeighInForm_x2_Horizontal-LETTER.xlsx",
)
REMOVED_ROWS = (8, 10, 13, 15)
EXPECTED_LEGACY_CELLS = {
    "A8": '<jx:if test="${ session.description != null && session.description != \'\' }">',
    "A9": "${ session.description }",
    "A10": "</jx:if>",
    "A13": '<jx:forEach items="${athletes}" var="l" varStatus="lifterLoop">',
    "A14": "${lifterLoop.index + 1}",
    "A15": "</jx:forEach>",
}
COMMENTS = {
    "A1": 'jx:area(lastCell="M15")',
    "A8": 'jx:if(condition="session.description != null && session.description != \'\'" lastCell="M8")',
    "A11": 'jx:each(items="athletes" var="l" varIndex="lifterIndex" lastCell="M11")',
}
EXPECTED_MERGES = {
    "B1:M1",
    "B5:L5",
    "L10:M10",
    "L11:M11",
    "C13:F13",
    "C14:F14",
    "C15:F15",
    "G13:H13",
    "G14:H14",
    "I13:M13",
    "I14:M14",
}
SIGNATURE_LABELS = {
    "B13": '${"1"}',
    "B14": '${"2"}',
    "B15": '${"3"}',
    "G13": '${t.get("CompetitionSecretary1")}',
    "G14": '${t.get("CompetitionSecretary2")}',
}


def remapped_row(row):
    return row - sum(removed < row for removed in REMOVED_ROWS)


def remove_and_remap_merges(sheet):
    ranges = list(sheet.merged_cells.ranges)
    for merged_range in ranges:
        sheet.unmerge_cells(str(merged_range))
    remapped = []
    for merged_range in ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        if any(min_row <= removed <= max_row for removed in REMOVED_ROWS):
            raise ValueError(f"Merged range {merged_range} intersects a command row")
        remapped.append((min_col, remapped_row(min_row), max_col, remapped_row(max_row)))
    return remapped


def restore_merges(sheet, ranges):
    for min_col, min_row, max_col, max_row in ranges:
        sheet.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row,
            end_column=max_col,
        )


def remap_row_dimensions(sheet):
    dimensions = [(row, copy.copy(dimension)) for row, dimension in sheet.row_dimensions.items()]
    sheet.row_dimensions.clear()
    for row, dimension in dimensions:
        if row in REMOVED_ROWS:
            continue
        new_row = remapped_row(row)
        dimension.index = new_row
        sheet.row_dimensions[new_row] = dimension


def normalize_header(sheet):
    expected = {
        "C1": '${t.get("Competition.competitionName")} :',
        "E1": "${competition.competitionName}",
        "C2": '${t.get("Competition.competitionSite")} :',
        "E2": "${competition.competitionSite} ${competition.competitionCity}",
        "C3": '${t.get("Competition.competitionDate")} :',
        "E3": "${competition.localizedCompetitionDate}",
    }
    mismatches = [
        f"{coordinate}: expected {value!r}, found {sheet[coordinate].value!r}"
        for coordinate, value in expected.items()
        if sheet[coordinate].value != value
    ]
    if mismatches:
        raise ValueError("Unexpected weigh-in header:\n  " + "\n  ".join(mismatches))

    label_styles = [copy.copy(sheet[f"C{row}"]._style) for row in range(1, 4)]
    value_styles = [copy.copy(sheet[f"E{row}"]._style) for row in range(1, 4)]
    values = [(sheet[f"C{row}"].value, sheet[f"E{row}"].value) for row in range(1, 4)]
    sheet.unmerge_cells("E1:M1")
    for row in range(1, 4):
        for column in range(1, 14):
            cell = sheet.cell(row, column)
            cell.value = None
            cell.hyperlink = None
            cell.comment = None
            cell.border = Border()
        label, value = values[row - 1]
        sheet.cell(row, 1, label)._style = label_styles[row - 1]
        sheet.cell(row, 2, value)._style = value_styles[row - 1]
        alignment = copy.copy(sheet.cell(row, 1).alignment)
        alignment.horizontal = "right"
        sheet.cell(row, 1).alignment = alignment
    sheet.merge_cells("B1:M1")
    for row in sheet.iter_rows(min_row=1, max_row=3, min_col=1, max_col=13):
        for cell in row:
            cell.border = Border()


def apply_signature_layout(sheet, require_legacy_footer):
    if require_legacy_footer:
        expected = {
            "C13": '${t.get("Weighin1")}',
            "E13": "${session.weighIn1}",
            "C14": '${t.get("Weighin2")}',
            "E14": "${session.weighIn2}",
        }
        mismatches = [
            f"{coordinate}: expected {value!r}, found {sheet[coordinate].value!r}"
            for coordinate, value in expected.items()
            if sheet[coordinate].value != value
        ]
        if mismatches:
            raise ValueError("Unexpected legacy signature footer:\n  " + "\n  ".join(mismatches))

    for merged_range in list(sheet.merged_cells.ranges):
        if merged_range.min_row >= 13:
            sheet.unmerge_cells(str(merged_range))
    line_style = copy.copy(sheet["E13"]._style if require_legacy_footer else sheet["C13"]._style)
    label_style = copy.copy(sheet["C13"]._style if require_legacy_footer else sheet["A13"]._style)
    row_height = sheet.row_dimensions[13].height
    for row in range(13, 16):
        for column in range(1, 14):
            cell = sheet.cell(row, column)
            cell.value = None
            cell.comment = None
            cell.hyperlink = None
            cell.border = Border()
        sheet.row_dimensions[row].height = row_height

    for coordinate, value in SIGNATURE_LABELS.items():
        sheet[coordinate] = value
        sheet[coordinate]._style = copy.copy(label_style)
        alignment = copy.copy(sheet[coordinate].alignment)
        alignment.horizontal = "center" if coordinate.startswith("B") else "right"
        sheet[coordinate].alignment = alignment
    for row in range(13, 16):
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        sheet.cell(row, 3)._style = copy.copy(line_style)
        sheet.cell(row, 3).border = Border(bottom=Side(style="thin", color="000000"))
    for row in range(13, 15):
        sheet.merge_cells(start_row=row, start_column=7, end_row=row, end_column=8)
        sheet.merge_cells(start_row=row, start_column=9, end_row=row, end_column=13)
        sheet.cell(row, 9)._style = copy.copy(line_style)
        sheet.cell(row, 9).border = Border(bottom=Side(style="thin", color="000000"))


def validate_legacy(sheet, path):
    mismatches = [
        f"{coordinate}: expected {value!r}, found {sheet[coordinate].value!r}"
        for coordinate, value in EXPECTED_LEGACY_CELLS.items()
        if sheet[coordinate].value != value
    ]
    if mismatches:
        raise ValueError(f"Unexpected legacy layout in {path}:\n  " + "\n  ".join(mismatches))


def is_converted(sheet):
    return sheet["A1"].comment is not None and sheet["A1"].comment.text == COMMENTS["A1"]


def has_previous_signature_layout(sheet):
    return (
        sheet["A1"].comment is not None
        and sheet["A1"].comment.text == 'jx:area(lastCell="M14")'
    )


def validate_converted(sheet, path):
    errors = []
    expected_values = {
        "A1": '${t.get("Competition.competitionName")} :',
        "B1": "${competition.competitionName}",
        "A2": '${t.get("Competition.competitionSite")} :',
        "B2": "${competition.competitionSite} ${competition.competitionCity}",
        "A3": '${t.get("Competition.competitionDate")} :',
        "B3": "${competition.localizedCompetitionDate}",
        "A8": "${ session.description }",
        "A10": '${t.get("Results.Start")}' if sheet["A10"].value is not None else None,
        "A11": "${lifterIndex + 1}",
        "D10": '${t.get("ScaleWeight")}',
        "L10": '${t.get("Jury.Signature")}',
        **SIGNATURE_LABELS,
    }
    for coordinate, value in expected_values.items():
        if sheet[coordinate].value != value:
            errors.append(f"{coordinate}: expected {value!r}, found {sheet[coordinate].value!r}")
    for coordinate, value in COMMENTS.items():
        actual = sheet[coordinate].comment.text if sheet[coordinate].comment else None
        if actual != value:
            errors.append(f"{coordinate} comment: expected {value!r}, found {actual!r}")
    markup = [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and ("<jx:" in cell.value or "</jx:" in cell.value)
    ]
    if markup:
        errors.append(f"legacy markup remains in {', '.join(markup)}")
    actual_merges = {str(merged_range) for merged_range in sheet.merged_cells.ranges}
    if actual_merges != EXPECTED_MERGES:
        errors.append(f"merged ranges: expected {sorted(EXPECTED_MERGES)}, found {sorted(actual_merges)}")
    for coordinate in ("A1", "A2", "A3"):
        if sheet[coordinate].alignment.horizontal != "right":
            errors.append(f"{coordinate} is not right-aligned")
    for coordinate in ("B13", "B14", "B15"):
        if sheet[coordinate].alignment.horizontal != "center":
            errors.append(f"{coordinate} signature number is not centered")
    for coordinate in ("G13", "G14"):
        if sheet[coordinate].alignment.horizontal != "right":
            errors.append(f"{coordinate} signature label is not right-aligned")
    for coordinate in ("C13", "C14", "C15", "I13", "I14"):
        if sheet[coordinate].border.bottom.style != "thin":
            errors.append(f"{coordinate} signature line is missing")
    bordered = [
        cell.coordinate
        for row in sheet.iter_rows(min_row=1, max_row=3, min_col=1, max_col=13)
        for cell in row
        if any(side is not None and side.style for side in (
            cell.border.left, cell.border.right, cell.border.top, cell.border.bottom))
    ]
    if bordered:
        errors.append(f"header borders remain in {', '.join(bordered)}")
    if errors:
        raise ValueError(f"Invalid converted layout in {path}:\n  " + "\n  ".join(errors))


def convert_sheet(sheet, path):
    validate_legacy(sheet, path)
    merges = remove_and_remap_merges(sheet)
    remap_row_dimensions(sheet)
    for row in sorted(REMOVED_ROWS, reverse=True):
        sheet.delete_rows(row)
    restore_merges(sheet, merges)
    normalize_header(sheet)
    sheet["A11"] = "${lifterIndex + 1}"
    apply_signature_layout(sheet, require_legacy_footer=True)
    for coordinate, command in COMMENTS.items():
        sheet[coordinate].comment = Comment(command, "owlcms")
    validate_converted(sheet, path)


def validate_workbook(workbook, path):
    for sheet in workbook.worksheets:
        validate_converted(sheet, path)


def convert(path, mode):
    workbook = load_workbook(path)
    if mode == "check":
        validate_workbook(workbook, path)
        print(f"Checked {path}")
        return
    if mode == "update":
        for sheet in workbook.worksheets:
            if has_previous_signature_layout(sheet):
                apply_signature_layout(sheet, require_legacy_footer=True)
                sheet["A1"].comment = Comment(COMMENTS["A1"], "owlcms")
            elif is_converted(sheet):
                apply_signature_layout(sheet, require_legacy_footer=False)
            else:
                raise ValueError(f"Template is not converted: {path}")
            validate_converted(sheet, path)
    elif any(is_converted(sheet) for sheet in workbook.worksheets):
        raise ValueError(f"Template is already converted: {path}")
    else:
        for sheet in workbook.worksheets:
            convert_sheet(sheet, path)
    if mode == "dry-run":
        print(f"Validated {path}")
        return
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        reopened = load_workbook(temporary_path)
        validate_workbook(reopened, temporary_path)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)
    print(f"Converted {path}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    modes = parser.add_mutually_exclusive_group(required=True)
    modes.add_argument("--dry-run", action="store_const", const="dry-run", dest="mode")
    modes.add_argument("--apply", action="store_const", const="apply", dest="mode")
    modes.add_argument("--update", action="store_const", const="update", dest="mode")
    modes.add_argument("--check", action="store_const", const="check", dest="mode")
    args = parser.parse_args()
    for name in TEMPLATE_NAMES:
        convert(TEMPLATE_DIRECTORY / name, args.mode)


if __name__ == "__main__":
    main()