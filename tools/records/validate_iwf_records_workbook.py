#!/usr/bin/env python3
"""Validate an owlCMS IWF records workbook against its scraped JSON."""
from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_LIFTS = {"Snatch", "Clean & Jerk", "Total"}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("json_input", type=Path)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()

    with args.json_input.open(encoding="utf-8") as input_file:
        records = json.load(input_file)

    expected_counts: dict[tuple[str, str], int] = defaultdict(int)
    for record in records:
        expected_counts[(record["AgeGroup"], record["M/F"])] += 1

    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    actual_counts: dict[tuple[str, str], int] = {}
    total_rows = 0

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        if not rows:
            raise ValueError(f"{worksheet.title}: sheet has no records")

        combinations = {(row[2], row[3]) for row in rows}
        if len(combinations) != 1:
            raise ValueError(f"{worksheet.title}: contains multiple age/gender combinations")

        combination = combinations.pop()
        categories: dict[tuple[object, object], set[str]] = defaultdict(set)
        for row in rows:
            categories[(row[6], row[7])].add(row[8])

        invalid_categories = {
            category: lifts
            for category, lifts in categories.items()
            if lifts != EXPECTED_LIFTS
        }
        if invalid_categories:
            raise ValueError(f"{worksheet.title}: invalid category lifts: {invalid_categories}")

        actual_counts[combination] = len(rows)
        total_rows += len(rows)
        print(f"{worksheet.title}: {len(rows)} records, {len(categories)} categories")

    if actual_counts != dict(expected_counts):
        raise ValueError(f"Workbook counts {actual_counts} do not match JSON counts {dict(expected_counts)}")
    if total_rows != len(records):
        raise ValueError(f"Workbook has {total_rows} records; JSON has {len(records)}")

    print(f"VALID: {total_rows} records across {len(workbook.sheetnames)} sheets")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())